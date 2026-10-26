"""Club ammo boxes (club/): shooters without accounts, stock from moves, roles. Synthetic data only."""
import sqlite3
import time

import pytest
from flask import Flask

import club
import diary
from diary import users


def _app():
    app = Flask(__name__)
    app.config["SECRET_KEY"] = "test"
    diary.init_app(app)
    club.init_app(app)
    return app


def _login(app, name, pw):
    c = app.test_client()
    r = c.post("/diary/login", data={"username": name, "password": pw})
    assert r.status_code == 302
    return c


@pytest.fixture()
def setup():
    """Club 'testclub': admin 'cadmin', holder 'chold' of 9 mm B; boxes 9 mm A/B, 7.62 A."""
    club.init_db()
    with sqlite3.connect(club.db_path()) as c:
        for t in ("clubs", "club_roles", "items", "boxes", "levels", "moves", "products", "sales", "matches", "competitors", "officer_counts"):
            c.execute(f"DELETE FROM {t}")
    club._posts.clear()
    for n in ("cadmin", "chold", "cshooter"):
        if not users.by_name(n):
            users.create(n, n + "-password-1")
    app = _app()
    cid, _slug = club.create_club("testclub", "Test Club")
    club.set_role(cid, users.by_name("cadmin")["id"], "admin")
    a = _login(app, "cadmin", "cadmin-password-1")
    a.put("/ammo/api/c/testclub", json={"name": "Test Club", "mobilepay": "12345", "pay_msg": "HARRASTEMAKSU"})
    ids = {}
    for cal, name, price in (("9 mm Luger", "Geco", 0.30), ("9 mm Luger", "S&B", 0.25), ("7.62x39", "", 0.40)):
        ids[name or "762"] = a.post("/ammo/api/c/testclub/items", json={"calibre": cal, "name": name, "price": price}).json["id"]
    for name, cal, extra in (("9 mm A", "9 mm Luger", {"low": 200, "type": "Pistooli"}),
                             ("9 mm B", "9 mm Luger", {"holder_user": "chold"}),
                             ("7.62 A", "7.62x39", {"type": "Kivääri"})):
        r = a.post("/ammo/api/c/testclub/boxes", json=dict(name=name, calibre=cal, **extra))
        assert r.status_code == 200, r.json
        ids[name] = r.json["id"]
    st = a.get("/ammo/api/c/testclub").json
    codes = {b["name"]: b["code"] for b in st["boxes"]}
    return app, a, ids, codes


def _today():
    return time.strftime("%Y-%m-%d")


def _stock(a, box):
    return next(b for b in a.get("/ammo/api/c/testclub").json["boxes"] if b["name"] == box)["total"]


def test_buy_use_move_count(setup):
    app, a, ids, codes = setup
    assert a.post("/ammo/api/c/testclub/moves", json={"kind": "buy", "box": ids["9 mm B"], "item": ids["Geco"], "qty": 1000}).status_code == 200
    assert a.post("/ammo/api/c/testclub/moves", json={"kind": "move", "box": ids["9 mm B"], "to": ids["9 mm A"], "item": ids["Geco"], "qty": 300}).status_code == 200
    anon = app.test_client()
    info = anon.get(f"/ammo/api/b/{codes['9 mm A']}").json
    assert info["box"]["calibre"] == "9 mm Luger" and info["club"]["mobilepay"] == "12345" and info["me"] is None
    assert info["club"]["pay_msg"] == "HARRASTEMAKSU"
    assert {i["name"] for i in info["items"]} == {"Geco"}                  # what's in the box
    assert {i["name"] for i in anon.get(f"/ammo/api/b/{codes['9 mm B']}").json["items"]} == {"Geco"}
    a.post("/ammo/api/c/testclub/moves", json={"kind": "buy", "box": ids["9 mm B"], "item": ids["S&B"], "qty": 50})
    assert {i["name"] for i in anon.get(f"/ammo/api/b/{codes['9 mm B']}").json["items"]} == {"Geco", "S&B"}
    assert [i["name"] for i in anon.get(f"/ammo/api/b/{codes['7.62 A']}").json["items"]] == [""]   # empty: every ammo of the calibre
    r = anon.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "Matti M", "date": _today(), "device": "dev1",
                                                               "client_id": "c1", "lines": [{"item": ids["Geco"], "qty": 150}]})
    assert r.status_code == 200 and r.json["total"] == 45.0 and len(r.json["ref"]) == 4
    # Sent again from the offline queue: no double
    r2 = anon.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "Matti M", "date": _today(), "device": "dev1",
                                                                "client_id": "c1", "lines": [{"item": ids["Geco"], "qty": 150}]})
    assert r2.json["ref"] == r.json["ref"]
    assert _stock(a, "9 mm A") == 150 and _stock(a, "9 mm B") == 750
    w = {x["kind"] + x.get("name", x.get("calibre", "")) for x in a.get("/ammo/api/c/testclub").json["warnings"]}
    assert "box9 mm A" in w                                                 # 150 < refill level 200
    a.post("/ammo/api/c/testclub/moves", json={"kind": "count", "box": ids["9 mm B"], "item": ids["Geco"], "qty": 690})
    assert _stock(a, "9 mm B") == 740
    # Ammo of another calibre can't go in, can't move more than there is
    assert a.post("/ammo/api/c/testclub/moves", json={"kind": "buy", "box": ids["9 mm A"], "item": ids["762"], "qty": 5}).status_code == 400
    assert a.post("/ammo/api/c/testclub/moves", json={"kind": "move", "box": ids["9 mm A"], "to": ids["9 mm B"], "item": ids["Geco"], "qty": 9999}).status_code == 400
    log = a.get("/ammo/api/c/testclub/log").json["moves"]
    assert any(m["kind"] == "use" and m["name"] == "Matti M" for m in log)
    assert all("device" not in m for m in log)


def test_shooter_validation_and_withdraw(setup):
    app, a, ids, codes = setup
    anon = app.test_client()
    url = f"/ammo/api/b/{codes['7.62 A']}/use"
    assert anon.post(url, json={"name": "M", "date": _today(), "lines": [{"item": ids["762"], "qty": 10}]}).status_code == 400
    assert anon.post(url, json={"name": "Matti", "date": "2001-01-01", "lines": [{"item": ids["762"], "qty": 10}]}).status_code == 400
    assert anon.post(url, json={"name": "Matti", "date": _today(), "lines": [{"item": ids["Geco"], "qty": 10}]}).status_code == 400
    assert anon.post(url, json={"name": "Matti", "date": _today(), "lines": [{"item": ids["762"], "qty": 0}]}).status_code == 400
    assert anon.get("/ammo/api/b/notarealcode123").status_code == 404
    r = anon.post(url, json={"name": "Matti", "date": _today(), "device": "devX", "paid": False,
                             "lines": [{"item": ids["762"], "qty": 30}]})
    grp = r.json["grp"]
    assert not r.json["paid"]
    mine = anon.post(f"/ammo/api/b/{codes['7.62 A']}/mine", json={"device": "devX"}).json["entries"]
    assert [e["grp"] for e in mine] == [grp]
    assert anon.post(f"/ammo/api/b/{codes['7.62 A']}/withdraw/{grp}", json={"device": "other"}).status_code == 404
    assert anon.post(f"/ammo/api/b/{codes['7.62 A']}/withdraw/{grp}", json={"device": "devX"}).status_code == 200
    assert _stock(a, "7.62 A") == 0
    # A new secret: the old link stops working
    old = codes["7.62 A"]
    a.post(f"/ammo/api/c/testclub/boxes/{ids['7.62 A']}/newcode", json={})
    assert anon.get(f"/ammo/api/b/{old}").status_code == 404


def test_roles(setup):
    app, a, ids, codes = setup
    anon, h = app.test_client(), _login(app, "chold", "chold-password-1")
    s = _login(app, "cshooter", "cshooter-password-1")
    assert anon.get("/ammo/api/c/testclub").status_code == 401
    assert s.get("/ammo/api/c/testclub").status_code == 404                 # no role: the club isn't there
    st = h.get("/ammo/api/c/testclub").json
    assert st["role"] == "holder"
    assert {b["name"]: bool(b["code"]) for b in st["boxes"]} == {"9 mm A": False, "9 mm B": True, "7.62 A": False}
    assert h.post("/ammo/api/c/testclub/moves", json={"kind": "buy", "box": ids["9 mm B"], "item": ids["S&B"], "qty": 500}).status_code == 200
    assert h.post("/ammo/api/c/testclub/moves", json={"kind": "buy", "box": ids["9 mm A"], "item": ids["S&B"], "qty": 500}).status_code == 403
    assert h.post("/ammo/api/c/testclub/items", json={"calibre": "x", "price": 1}).status_code == 403
    anon.post(f"/ammo/api/b/{codes['9 mm B']}/use", json={"name": "Secret Name", "date": _today(), "lines": [{"item": ids["S&B"], "qty": 10}]})
    assert all(m["name"] != "Secret Name" for m in h.get("/ammo/api/c/testclub/log").json["moves"])
    assert h.get("/ammo/testclub/admin").status_code == 200 and s.get("/ammo/testclub/admin").status_code == 404
    assert a.get("/ammo/testclub/qr").status_code == 200 and h.get("/ammo/testclub/qr").status_code == 404


def test_logged_in_fills_name_and_adds_to_diary(setup):
    app, a, ids, codes = setup
    s = _login(app, "cshooter", "cshooter-password-1")
    s.put("/diary/api/profile", json={"name": "Seppo Shooter"})
    assert s.get(f"/ammo/api/b/{codes['9 mm A']}").json["me"]["name"] == "Seppo Shooter"
    r = s.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "Seppo Shooter", "date": _today(), "diary": True,
                                                            "device": "d", "lines": [{"item": ids["Geco"], "qty": 50}]})
    assert r.json["diary"] is True
    visits = s.get("/diary/api/sessions").json
    line = [ln for v in visits if v["date"] == _today() for ln in v["lines"] if ln.get("club") == r.json["grp"]]
    assert line and line[0]["rounds"] == 50 and line[0]["type"] == "Pistooli" and line[0]["cal"] == "9 mm Luger"
    s.post(f"/ammo/api/b/{codes['9 mm A']}/withdraw/{r.json['grp']}", json={"device": "d"})
    assert not [v for v in s.get("/diary/api/sessions").json if any(ln.get("club") == r.json["grp"] for ln in v["lines"])]


def test_bridge_is_the_only_diary_import():
    import pathlib
    for p in pathlib.Path(club.__file__).parent.glob("*.py"):
        if p.name != "bridge.py":
            assert "diary" not in [ln.split()[1].split(".")[0] for ln in p.read_text().splitlines()
                                   if ln.startswith(("import ", "from "))], p


def test_club_id_from_name_and_delete():
    club.init_db()
    with club.connect() as c:
        c.execute("DELETE FROM clubs WHERE slug LIKE 'vaares%'")
    assert club.create_club("", "VaaRes SRA")[1] == "vaares-sra"
    assert club.create_club("", "Vaares  SRA!")[1] == "vaares-sra-2"
    assert club.create_club("", "Ähtärin Ampujat")[1] == "ahtarin-ampujat"
    from club.__main__ import main
    with pytest.raises(SystemExit):
        main(["club-delete", "nosuchclub"])
    main(["club-delete", "vaares-sra-2"])
    assert club.club_by_slug("vaares-sra-2", club.connect()) is None
    assert club.club_by_slug("vaares-sra", club.connect())


def test_box_mobilepay_link(setup):
    app, a, ids, codes = setup
    st = a.get("/ammo/api/c/testclub").json
    b = next(x for x in st["boxes"] if x["name"] == "7.62 A")
    body = {k: b[k] for k in ("id", "name", "calibre", "type", "low", "holder_name")}
    assert a.post("/ammo/api/c/testclub/boxes", json=dict(body, pay_url="javascript:alert(1)")).status_code == 400
    assert a.post("/ammo/api/c/testclub/boxes", json=dict(body, pay_url="https://qr.mobilepay.fi/example")).status_code == 200
    assert app.test_client().get(f"/ammo/api/b/{codes['7.62 A']}").json["box"]["pay_url"] == "https://qr.mobilepay.fi/example"


def test_fees_and_merch(setup):
    app, a, ids, codes = setup
    P = "/ammo/api/c/testclub/products"
    season = a.post(P, json={"kind": "season", "name": "Kausimaksu", "price": 40, "pay_msg": "KAUSIMAKSU",
                             "pay_url": "https://qr.mobilepay.fi/vp/season"}).json["id"]
    single = a.post(P, json={"kind": "single", "name": "Kertamaksu", "price": 10, "pay_msg": "KERTAMAKSU",
                             "pay_url": "https://qr.mobilepay.fi/vp/single"}).json["id"]
    patch = a.post(P, json={"kind": "merch", "name": "Hihamerkki", "price": 8, "pay_msg": "KERTAMAKSU",
                            "pay_url": "https://qr.mobilepay.fi/vp/single", "track": True, "low": 3}).json["id"]
    assert a.post(P, json={"kind": "bogus", "name": "x", "price": 1}).status_code == 400
    st = a.get("/ammo/api/c/testclub").json
    code = st["pay_code"]
    anon = app.test_client()
    # Merch without stock isn't offered
    assert {p["name"] for p in anon.get(f"/ammo/api/p/{code}").json["products"]} == {"Kausimaksu", "Kertamaksu"}
    a.post("/ammo/api/c/testclub/pmoves", json={"kind": "refill", "product": patch, "qty": 5})
    assert "Hihamerkki" in {p["name"] for p in anon.get(f"/ammo/api/p/{code}").json["products"]}
    assert anon.post(f"/ammo/api/p/{code}/pay", json={"name": "Matti M", "date": _today(),
                                                      "lines": [{"product": season}, {"product": single}]}).status_code == 400
    r = anon.post(f"/ammo/api/p/{code}/pay", json={"name": "Matti M", "date": _today(), "device": "d1", "client_id": "p1",
                                                   "lines": [{"product": season}]}).json
    assert r["total"] == 40.0 and r["pays"][0]["msg"] == "KAUSIMAKSU " + r["ref"]
    r2 = anon.post(f"/ammo/api/p/{code}/pay", json={"name": "Matti M", "date": _today(), "device": "d1",
                                                    "lines": [{"product": patch, "qty": 3}]}).json
    assert r2["total"] == 24.0
    # Resent from the offline queue: no double; can't buy more patches than there are
    assert anon.post(f"/ammo/api/p/{code}/pay", json={"name": "Matti M", "date": _today(), "client_id": "p1",
                                                      "lines": [{"product": single}]}).json["ref"] == r["ref"]
    assert anon.post(f"/ammo/api/p/{code}/pay", json={"name": "X Y", "date": _today(),
                                                      "lines": [{"product": patch, "qty": 3}]}).status_code == 400
    st = a.get("/ammo/api/c/testclub").json
    assert next(p for p in st["products"] if p["id"] == patch)["stock"] == 2
    assert any(w["kind"] == "merch" for w in st["warnings"])                   # 2 < 3
    a.post("/ammo/api/c/testclub/pmoves", json={"kind": "count", "product": patch, "qty": 4})
    assert next(p for p in a.get("/ammo/api/c/testclub").json["products"] if p["id"] == patch)["stock"] == 4
    # Season list, log, void
    assert [x["name"] for x in a.get("/ammo/api/c/testclub/season").json["rows"]] == ["Matti M"]
    log = a.get("/ammo/api/c/testclub/log").json["moves"]
    assert {m["kind"] for m in log} >= {"sale", "prefill", "pcount"}
    assert a.post(f"/ammo/api/c/testclub/paid/{r['grp']}", json={"paid": False}).json["ok"]
    a.post(f"/ammo/api/c/testclub/void/{r['grp']}", json={})
    assert a.get("/ammo/api/c/testclub/season").json["rows"] == []
    # Box holders don't see fees; a new link stops the old one
    h = _login(app, "chold", "chold-password-1")
    assert h.get("/ammo/api/c/testclub").json["pay_code"] is None
    assert all(m["kind"] not in ("sale", "prefill", "pcount") for m in h.get("/ammo/api/c/testclub/log").json["moves"])
    a.post("/ammo/api/c/testclub/paycode", json={})
    assert anon.get(f"/ammo/api/p/{code}").status_code == 404


def test_pay_message_template():
    from club.shop import pay_message
    assert pay_message("SRA KERTAMAKSU {pvm}", "2026-10-27", 1, "A7K3") == "SRA KERTAMAKSU 27.10.2026 A7K3"
    assert pay_message("SRA KAUSIMAKSU {vuosi}", "2026-10-27", 1, "A7K3") == "SRA KAUSIMAKSU 2026 A7K3"
    assert pay_message("SRA Hihamerkki {kpl} kpl", "2026-10-27", 2, "A7K3") == "SRA Hihamerkki 2 kpl A7K3"
    assert pay_message("{koodi} SRA", "2026-10-27", 1, "A7K3") == "A7K3 SRA"
    assert pay_message("", "2026-10-27", 1, "A7K3") == "A7K3"


def test_merch_quantity_price(setup):
    app, a, ids, codes = setup
    from club.shop import line_price
    p = {"price": 10, "deal_qty": 3, "deal_price": 25}
    assert [line_price(p, q) for q in range(1, 7)] == [10, 20, 25, 35, 45, 50]
    pid = a.post("/ammo/api/c/testclub/products", json={"kind": "merch", "name": "SRA Hihamerkki", "price": 10, "deal_qty": 3,
                                                        "deal_price": 25, "pay_msg": "SRA Hihamerkki {kpl} kpl",
                                                        "track": True}).json["id"]
    a.post("/ammo/api/c/testclub/pmoves", json={"kind": "refill", "product": pid, "qty": 10})
    code = a.get("/ammo/api/c/testclub").json["pay_code"]
    r = app.test_client().post(f"/ammo/api/p/{code}/pay", json={"name": "Matti M", "date": _today(),
                                                                "lines": [{"product": pid, "qty": 4}]}).json
    assert r["total"] == 35.0 and r["pays"][0]["msg"] == f"SRA Hihamerkki 4 kpl {r['ref']}"
    st = a.get("/ammo/api/c/testclub").json
    assert next(p for p in st["products"] if p["id"] == pid)["stock"] == 6
    assert next(m for m in a.get("/ammo/api/c/testclub/log").json["moves"] if m["kind"] == "sale")["amount"] == 35.0


def test_product_delete_only_without_history(setup):
    app, a, ids, codes = setup
    P = "/ammo/api/c/testclub/products"
    unused = a.post(P, json={"kind": "merch", "name": "Old pack", "price": 25, "archived": True}).json["id"]
    used = a.post(P, json={"kind": "single", "name": "Kertamaksu", "price": 10}).json["id"]
    code = a.get("/ammo/api/c/testclub").json["pay_code"]
    app.test_client().post(f"/ammo/api/p/{code}/pay", json={"name": "Matti M", "date": _today(), "lines": [{"product": used}]})
    assert a.post(f"{P}/{used}/delete", json={}).status_code == 400
    assert a.post(f"{P}/{unused}/delete", json={}).status_code == 200
    assert {p["id"] for p in a.get("/ammo/api/c/testclub").json["products"]} == {used}


def test_box_and_item_delete_only_without_history(setup):
    app, a, ids, codes = setup
    C = "/ammo/api/c/testclub"
    # 7.62 A: a hand-over only (no rounds ever) → can go; its unused ammo price too
    assert a.post(f"{C}/boxes/{ids['7.62 A']}/delete", json={}).status_code == 200
    assert a.post(f"{C}/items/{ids['762']}/delete", json={}).status_code == 200
    # 9 mm B gets rounds → stays; the ammo used → stays
    a.post(f"{C}/moves", json={"kind": "buy", "box": ids["9 mm B"], "item": ids["Geco"], "qty": 10})
    r = a.post(f"{C}/boxes/{ids['9 mm B']}/delete", json={})
    assert r.status_code == 400 and r.json["error"] == "in_use"
    assert a.post(f"{C}/items/{ids['Geco']}/delete", json={}).status_code == 400
    st = a.get(C).json
    assert {b["name"] for b in st["boxes"]} == {"9 mm A", "9 mm B"} and "" not in {i["name"] for i in st["items"]}


def test_treasurer_export(setup):
    app, a, ids, codes = setup
    import io
    from openpyxl import load_workbook
    anon = app.test_client()
    a.post("/ammo/api/c/testclub/moves", json={"kind": "buy", "box": ids["9 mm A"], "item": ids["Geco"], "qty": 500})
    anon.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "Matti Meikäläinen", "date": _today(),
                                                          "lines": [{"item": ids["Geco"], "qty": 50}]})
    anon.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "Liisa L", "date": _today(), "paid": False,
                                                          "lines": [{"item": ids["Geco"], "qty": 10}]})
    fee = a.post("/ammo/api/c/testclub/products", json={"kind": "single", "name": "Kertamaksu", "price": 10}).json["id"]
    code = a.get("/ammo/api/c/testclub").json["pay_code"]
    anon.post(f"/ammo/api/p/{code}/pay", json={"name": "Matti Meikäläinen", "date": _today(), "lines": [{"product": fee}]})
    E = f"/ammo/api/c/testclub/export?from={_today()}&to={_today()}"
    j = a.get(E + "&fmt=json").json
    assert j["total"] == {"count": 3, "amount": 28.0, "paid": 25.0, "unpaid": 3.0}
    assert {(s["kind"], s["what"], s["qty"]) for s in j["summary"]} == {("ammo", "9 mm Luger · Geco", 60), ("single", "Kertamaksu", 1)}
    assert a.get(E + "&fmt=json&what=fees").json["total"]["count"] == 1
    assert a.get(E + "&fmt=json&unpaid=1").json["total"]["amount"] == 3.0
    c = a.get(E + "&fmt=csv&lang=fi")
    text = c.get_data(as_text=True)
    assert c.headers["Content-Disposition"].startswith("attachment") and text.startswith("﻿Päivämäärä;Nimi;")
    assert "Matti Meikäläinen;Patruunat;9 mm Luger · Geco;9 mm A;50;15,00;kyllä;" in text
    x = a.get(E + "&fmt=xlsx&lang=fi")
    wb = load_workbook(io.BytesIO(x.data))
    assert wb.sheetnames == ["Tapahtumat", "Yhteenveto"] and wb["Tapahtumat"].max_row == 4
    h = _login(app, "chold", "chold-password-1")
    assert h.get(E + "&fmt=csv").status_code == 403                    # box holders: no names, no money


def test_box_with_two_calibres(setup):
    app, a, ids, codes = setup
    C = "/ammo/api/c/testclub"
    sg = a.post(f"{C}/items", json={"calibre": "12/70", "name": "Trap", "price": 0.40}).json["id"]
    rf = a.post(f"{C}/items", json={"calibre": "7.62x54R", "name": "", "price": 0.80}).json["id"]
    bid = a.post(f"{C}/boxes", json={"name": "Charlie", "calibre": "12/70", "cals": ["7.62x54R"], "type": ""}).json["id"]
    a.post(f"{C}/moves", json={"kind": "buy", "box": bid, "item": sg, "qty": 100})
    a.post(f"{C}/moves", json={"kind": "buy", "box": bid, "item": rf, "qty": 60})
    assert a.post(f"{C}/moves", json={"kind": "buy", "box": bid, "item": ids["Geco"], "qty": 5}).status_code == 400
    st = a.get(C).json
    box = next(b for b in st["boxes"] if b["id"] == bid)
    assert box["cals"] == ["12/70", "7.62x54R"] and box["total"] == 160
    assert st["cal_totals"]["12/70"] == 100 and st["cal_totals"]["7.62x54R"] == 60
    info = app.test_client().get(f"/ammo/api/b/{box['code']}").json
    assert info["box"]["calibre"] == "12/70 + 7.62x54R" and [i["calibre"] for i in info["items"]] == ["12/70", "7.62x54R"]
    s = _login(app, "cshooter", "cshooter-password-1")
    r = s.post(f"/ammo/api/b/{box['code']}/use", json={"name": "Seppo", "date": _today(), "diary": True,
                                                       "lines": [{"item": sg, "qty": 25}, {"item": rf, "qty": 10}]}).json
    assert r["total"] == 18.0 and r["diary"] is True
    lines = [ln for v in s.get("/diary/api/sessions").json for ln in v["lines"] if ln.get("club") == r["grp"]]
    assert sorted((ln["cal"], ln["rounds"]) for ln in lines) == [("12/70", 25), ("7.62x54R", 10)]
    # Can't drop a calibre that still has rounds in the box
    assert a.post(f"{C}/boxes", json={"id": bid, "name": "Charlie", "calibre": "12/70", "cals": []}).status_code == 400


def test_refill_level_per_calibre(setup):
    app, a, ids, codes = setup
    C = "/ammo/api/c/testclub"
    sg = a.post(f"{C}/items", json={"calibre": "12/70", "name": "", "price": 0.40}).json["id"]
    rf = a.post(f"{C}/items", json={"calibre": "7.62x54R", "name": "", "price": 0.80}).json["id"]
    bid = a.post(f"{C}/boxes", json={"name": "Charlie", "calibre": "12/70", "cals": ["7.62x54R"],
                                     "lows": {"12/70": 50, "7.62x54R": 40}}).json["id"]
    a.post(f"{C}/moves", json={"kind": "buy", "box": bid, "item": sg, "qty": 100})
    a.post(f"{C}/moves", json={"kind": "buy", "box": bid, "item": rf, "qty": 30})
    st = a.get(C).json
    assert next(b for b in st["boxes"] if b["id"] == bid)["lows"] == {"12/70": 50, "7.62x54R": 40}
    w = [x for x in st["warnings"] if x["kind"] == "box" and x["box"] == bid]
    assert [(x["calibre"], x["n"], x["low"]) for x in w] == [("7.62x54R", 30, 40)]      # 12/70: 100 ≥ 50
    # An old single level stays on the box's main calibre
    old = next(b for b in st["boxes"] if b["name"] == "9 mm A")
    assert old["lows"] == {"9 mm Luger": 200}


def test_match_officer_counts_and_report(setup):
    app, a, ids, codes = setup
    C = "/ammo/api/c/testclub"
    if not users.by_name("cofficer"):
        users.create("cofficer", "cofficer-password-1")
    assert a.post(f"{C}/people", json={"username": "cofficer", "role": "officer"}).status_code == 200
    a.post(f"{C}/moves", json={"kind": "buy", "box": ids["9 mm A"], "item": ids["Geco"], "qty": 1000})
    mid = a.post(f"{C}/matches", json={"name": "Test match", "date": _today(), "stages": 4,
                                       "boxes": [ids["9 mm A"]]}).json["id"]
    r = a.post(f"{C}/matches/{mid}/competitors", json={"replace": True, "rows": [
        {"num": "1", "name": "Aatu Ampuja", "squad": "1", "division": "Production"},
        {"num": "2", "name": "Bertta Bee", "squad": "1"}, {"num": "3", "name": "Cecilia Cee", "squad": "2"},
        {"name": "Aatu Ampuja"}, {"name": "x"}]}).json
    assert r["added"] == 3
    st = a.get(f"{C}/matches/{mid}").json
    comp = {c["name"]: c["id"] for c in st["competitors"]}
    assert [i["name"] for i in st["items"]] == ["Geco", "S&B"]          # the 9 mm ammo
    # The officer: only the officer page and its API
    o = _login(app, "cofficer", "cofficer-password-1")
    assert o.get("/ammo/testclub/admin").status_code == 302 and o.get("/ammo/testclub/ro").status_code == 200
    assert o.get(C).status_code == 403 and o.get(f"{C}/matches").status_code == 403
    assert o.get(f"{C}/ro").json["matches"][0]["id"] == mid
    for cname, stage, q in (("Aatu Ampuja", 1, 30), ("Aatu Ampuja", 2, 20), ("Bertta Bee", 1, 25), ("Cecilia Cee", 3, 40)):
        assert o.post(f"{C}/matches/{mid}/count", json={"comp": comp[cname], "stage": stage, "item": ids["Geco"],
                                                        "qty": q}).status_code == 200
    assert o.post(f"{C}/matches/{mid}/count", json={"comp": comp["Aatu Ampuja"], "stage": 9, "item": ids["Geco"],
                                                    "qty": 5}).status_code == 400            # no stage 9
    assert "report" not in o.get(f"{C}/matches/{mid}").json                                  # officers: no money
    # Shooters: the box form offers the competitor list; one picks their name, one types it
    anon = app.test_client()
    info = anon.get(f"/ammo/api/b/{codes['9 mm A']}").json
    assert info["match"]["id"] == mid and len(info["match"]["competitors"]) == 3
    anon.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "", "comp": comp["Aatu Ampuja"], "date": _today(),
                                                          "lines": [{"item": ids["Geco"], "qty": 50}]})
    anon.post(f"/ammo/api/b/{codes['9 mm A']}/use", json={"name": "bertta bee", "date": _today(), "paid": False,
                                                          "lines": [{"item": ids["Geco"], "qty": 20}]})
    rep = {x["name"]: x for x in a.get(f"{C}/matches/{mid}").json["report"]}
    assert (rep["Aatu Ampuja"]["officer"], rep["Aatu Ampuja"]["own"], rep["Aatu Ampuja"]["status"]) == (50, 50, "ok")
    assert (rep["Bertta Bee"]["officer"], rep["Bertta Bee"]["own"], rep["Bertta Bee"]["status"]) == (25, 20, "less")
    assert rep["Bertta Bee"]["owes"] == 1.5
    assert (rep["Cecilia Cee"]["own"], rep["Cecilia Cee"]["status"]) == (0, "missing")
    # 0 removes an officer entry
    o.post(f"{C}/matches/{mid}/count", json={"comp": comp["Cecilia Cee"], "stage": 3, "item": ids["Geco"], "qty": 0})
    assert "Cecilia Cee" not in {x["name"] for x in a.get(f"{C}/matches/{mid}").json["report"]}
    # Default ammo per stage: kept for real stages and ammo ids, stage 9 of 4 dropped
    a.post(f"{C}/matches", json={"id": mid, "name": "Test match", "date": _today(), "stages": 4, "boxes": [ids["9 mm A"]],
                                 "stage_ammo": {"1": [ids["Geco"]], "2": [ids["Geco"], ids["S&B"], ids["Geco"]], "9": [ids["S&B"]]}})
    assert o.get(f"{C}/matches/{mid}").json["match"]["stage_ammo"] == {"1": [ids["Geco"]], "2": [ids["Geco"], ids["S&B"]]}
    # Two ammo types on one stage (e.g. pistol + rifle): both kept, both in the report
    o.post(f"{C}/matches/{mid}/count", json={"comp": comp["Cecilia Cee"], "stage": 2, "item": ids["Geco"], "qty": 10})
    o.post(f"{C}/matches/{mid}/count", json={"comp": comp["Cecilia Cee"], "stage": 2, "item": ids["S&B"], "qty": 15})
    cc = {x["name"]: x for x in a.get(f"{C}/matches/{mid}").json["report"]}["Cecilia Cee"]
    assert cc["officer"] == 25 and {ln["item"]: ln["officer"] for ln in cc["lines"]} == {ids["Geco"]: 10, ids["S&B"]: 15}
    # The list can't be replaced once officers have entered rounds
    assert a.post(f"{C}/matches/{mid}/competitors", json={"replace": True, "rows": []}).status_code == 400



SSI_EVENT = {"event": {"id": "1", "name": "Testikisa 2026", "starts": "2026-10-27T15:00:00+00:00", "stages_count": 5,
                       "competitors_approved_w_wo_results_not_dnf": [
                           {"id": "c1", "first_name": "Aatu", "last_name": "Ampuja", "number": 1, "get_division_display": "Production"},
                           {"id": "c2", "first_name": "Bertta", "last_name": "Bee", "number": 2, "get_division_display": "Standard"},
                           {"id": "c3", "first_name": "Cecilia", "last_name": "Cee", "number": 3, "get_division_display": "Open"}],
                       "squads": [{"id": "s1", "number": 1, "competitors": [{"id": "c1"}, {"id": "c2"}]},
                                  {"id": "s2", "number": 2, "competitors": [{"id": "c3"}]}]}}


def _fake_ssi(monkeypatch, expire_once=False):
    from club import ssi
    for k, v in (("SSI_API_KEY", "test-key"), ("SSI_EMAIL", "svc@example.com"), ("SSI_PASSWORD", "x")):
        monkeypatch.setenv(k, v)
    ssi._token.update(jwt="", at=0)
    calls = {"login": 0, "query": 0, "expired": expire_once}

    def post(q, variables, jwt=None, opener=None):
        if "token_auth" in q:
            calls["login"] += 1
            return {"token_auth": {"success": True, "token": {"token": "jwt-%d" % calls["login"]}}}
        assert jwt and variables == {"ct": 22, "id": "1"}
        calls["query"] += 1
        if calls["expired"]:
            calls["expired"] = False
            raise ssi.SsiError("SSI: Signature has expired")
        return SSI_EVENT
    monkeypatch.setattr(ssi, "_post", post)
    return calls


def test_ssi_api_reader(monkeypatch):
    from club import ssi
    assert ssi.parse_url("https://shootnscoreit.com/event/22/29338/participants/") == ("22", "29338")
    assert ssi.parse_url("https://evil.example/event/22/1/") is None
    monkeypatch.delenv("SSI_API_KEY", raising=False)
    with pytest.raises(ssi.SsiError):
        ssi.fetch("https://shootnscoreit.com/event/22/1/")
    calls = _fake_ssi(monkeypatch, expire_once=True)
    r = ssi.fetch("https://shootnscoreit.com/event/22/1/participants/")
    assert (r["name"], r["date"], r["stages"]) == ("Testikisa 2026", "2026-10-27", 5)
    assert r["competitors"][0] == {"num": "1", "name": "Aatu Ampuja", "squad": "1", "division": "Production"}
    assert [c["squad"] for c in r["competitors"]] == ["1", "1", "2"]
    assert calls == {"login": 2, "query": 2, "expired": False}            # expired token: logged in again once
    ssi.fetch("https://shootnscoreit.com/event/22/1/")
    assert calls["login"] == 2                                             # token kept in memory


def test_ssi_endpoint_and_match_link(setup, monkeypatch):
    app, a, _ids, _codes = setup
    C = "/ammo/api/c/testclub"
    assert a.get(f"{C}/matches").json["ssi"] in (True, False)
    _fake_ssi(monkeypatch)
    assert a.get(f"{C}/matches").json["ssi"] is True
    r = a.post(f"{C}/ssi", json={"url": "https://shootnscoreit.com/event/22/1/"}).json
    assert r["name"] == "Testikisa 2026" and len(r["competitors"]) == 3
    assert a.post(f"{C}/ssi", json={"url": "https://example.com/"}).status_code == 400
    mid = a.post(f"{C}/matches", json={"name": r["name"], "date": r["date"], "stages": r["stages"],
                                       "ssi_url": "https://shootnscoreit.com/event/22/1/participants/?x=1"}).json["id"]
    assert a.get(f"{C}/matches/{mid}").json["match"]["ssi_url"] == "https://shootnscoreit.com/event/22/1/"
    assert a.post(f"{C}/matches/{mid}/competitors", json={"rows": r["competitors"]}).json["added"] == 3
    moved = [dict(c, squad="3") if c["name"] == "Cecilia Cee" else c for c in r["competitors"]]   # re-squadded in SSI
    c = a.post(f"{C}/matches/{mid}/competitors", json={"rows": moved}).json
    assert (c["added"], c["updated"]) == (0, 1)
    assert {x["name"]: x["squad"] for x in a.get(f"{C}/matches/{mid}").json["competitors"]}["Cecilia Cee"] == "3"
    assert app.test_client().post(f"{C}/ssi", json={"url": "https://shootnscoreit.com/event/22/1/"}).status_code in (401, 403)
