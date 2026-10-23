"""Treasurer's export (admins): the payments of a period – ammo taken from the boxes and fees / merch paid – as CSV
(Finnish Excel: ; and decimal comma), Excel (entries + summary sheets) or JSON for the PDF the admin page makes.
Voided entries are left out. One row per ammo / product line.

  GET /ammo/api/c/<club>/export?from=YYYY-MM-DD&to=YYYY-MM-DD&what=all|ammo|fees&unpaid=1&lang=fi|en&fmt=csv|xlsx|json
"""
import csv
import io
import re

from flask import Response, abort, jsonify, request

from . import DATE_RE, _club_for, cbp, cdb

COLS = {
    "fi": ["Päivämäärä", "Nimi", "Tyyppi", "Mitä", "Laatikko", "Määrä", "Summa €", "Maksettu", "Koodi"],
    "en": ["Date", "Name", "Type", "What", "Box", "Qty", "Amount €", "Paid", "Code"],
}
WORDS = {
    "fi": {"ammo": "Patruunat", "season": "Kausimaksu", "single": "Kertamaksu", "merch": "Tuote", "yes": "kyllä", "no": "ei",
           "entries": "Tapahtumat", "summary": "Yhteenveto", "total": "Yhteensä", "paid": "Maksettu", "unpaid": "Maksamatta",
           "count": "Kirjauksia", "sum_cols": ["Tyyppi", "Mitä", "Määrä", "Summa €", "Maksettu €", "Maksamatta €"]},
    "en": {"ammo": "Ammo", "season": "Season fee", "single": "One-time fee", "merch": "Merch", "yes": "yes", "no": "no",
           "entries": "Entries", "summary": "Summary", "total": "Total", "paid": "Paid", "unpaid": "Not paid",
           "count": "Entries", "sum_cols": ["Type", "What", "Qty", "Amount €", "Paid €", "Not paid €"]},
}


def rows(conn, club_id, start, end, what, unpaid_only):
    out = []
    if what in ("all", "ammo"):
        for r in conn.execute(
                "SELECT m.date, m.name, m.qty, m.price, m.paid, m.ref, m.at, i.calibre, i.name AS iname, b.name AS bname "
                "FROM moves m LEFT JOIN items i ON i.id=m.item_id LEFT JOIN boxes b ON b.id=m.box_id "
                "WHERE m.club_id=? AND m.kind='use' AND m.voided_at IS NULL AND m.date BETWEEN ? AND ?",
                (club_id, start, end)):
            out.append({"date": r["date"], "at": r["at"], "name": r["name"], "kind": "ammo",
                        "what": " · ".join(x for x in (r["calibre"], r["iname"]) if x), "box": r["bname"] or "",
                        "qty": -r["qty"], "amount": round(-r["qty"] * (r["price"] or 0), 2), "paid": bool(r["paid"]),
                        "ref": r["ref"] or ""})
    if what in ("all", "fees"):
        for r in conn.execute(
                "SELECT s.date, s.name, s.qty, s.price, s.amount, s.paid, s.ref, s.at, p.name AS pname, p.kind AS pkind "
                "FROM sales s JOIN products p ON p.id=s.product_id "
                "WHERE s.club_id=? AND s.kind='sale' AND s.voided_at IS NULL AND s.date BETWEEN ? AND ?",
                (club_id, start, end)):
            amount = r["amount"] if r["amount"] is not None else round(r["qty"] * (r["price"] or 0), 2)
            out.append({"date": r["date"], "at": r["at"], "name": r["name"], "kind": r["pkind"], "what": r["pname"],
                        "box": "", "qty": r["qty"], "amount": amount, "paid": bool(r["paid"]), "ref": r["ref"] or ""})
    if unpaid_only:
        out = [x for x in out if not x["paid"]]
    out.sort(key=lambda x: (x["date"], x["at"]))
    return out


def summary(entries):
    by = {}
    for e in entries:
        s = by.setdefault((e["kind"], e["what"]), {"kind": e["kind"], "what": e["what"], "qty": 0, "amount": 0.0,
                                                   "paid": 0.0, "unpaid": 0.0})
        s["qty"] += e["qty"]
        s["amount"] = round(s["amount"] + e["amount"], 2)
        s["paid" if e["paid"] else "unpaid"] = round(s["paid" if e["paid"] else "unpaid"] + e["amount"], 2)
    order = {"ammo": 0, "season": 1, "single": 2, "merch": 3}
    lines = sorted(by.values(), key=lambda s: (order.get(s["kind"], 9), s["what"]))
    tot = {"count": len(entries), "amount": round(sum(e["amount"] for e in entries), 2),
           "paid": round(sum(e["amount"] for e in entries if e["paid"]), 2),
           "unpaid": round(sum(e["amount"] for e in entries if not e["paid"]), 2)}
    return lines, tot


def _fi_date(d):
    y, m, dd = d.split("-")
    return f"{int(dd)}.{int(m)}.{y}"


@cbp.get("/api/c/<slug>/export")
def api_export(slug):
    club, _u, _r = _club_for(slug)
    a = request.args
    start, end = a.get("from") or "0000-01-01", a.get("to") or "9999-12-31"
    if not (DATE_RE.match(start) and DATE_RE.match(end)) or start > end:
        abort(400)
    what = a.get("what") if a.get("what") in ("all", "ammo", "fees") else "all"
    lang = "fi" if a.get("lang") == "fi" else "en"
    fmt = a.get("fmt") if a.get("fmt") in ("csv", "xlsx", "json") else "json"
    entries = rows(cdb(), club["id"], start, end, what, a.get("unpaid") == "1")
    lines, tot = summary(entries)
    W = WORDS[lang]
    kind = lambda k: W.get(k, k)                                        # noqa: E731
    base = re.sub(r"[^a-z0-9-]+", "-", f"{club['slug']}-{start}-{end}").strip("-")

    if fmt == "json":
        return jsonify(club=club["name"], start=start, end=end, what=what, entries=entries, summary=lines, total=tot)

    if fmt == "csv":
        buf = io.StringIO()
        buf.write("﻿")                                            # so Excel reads it as UTF-8
        w = csv.writer(buf, delimiter=";" if lang == "fi" else ",", lineterminator="\r\n")
        w.writerow(COLS[lang])
        num = (lambda x: f"{x:.2f}".replace(".", ",")) if lang == "fi" else (lambda x: f"{x:.2f}")
        for e in entries:
            w.writerow([_fi_date(e["date"]) if lang == "fi" else e["date"], e["name"], kind(e["kind"]), e["what"], e["box"],
                        e["qty"], num(e["amount"]), W["yes"] if e["paid"] else W["no"], e["ref"]])
        return Response(buf.getvalue(), mimetype="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{base}.csv"', "Cache-Control": "no-store"})

    # Excel: entries + summary, real dates and numbers
    from datetime import date
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = W["entries"]
    ws.append(COLS[lang])
    for c in ws[1]:
        c.font = Font(bold=True)
    for e in entries:
        ws.append([date.fromisoformat(e["date"]), e["name"], kind(e["kind"]), e["what"], e["box"], e["qty"], e["amount"],
                   W["yes"] if e["paid"] else W["no"], e["ref"]])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "d.m.yyyy" if lang == "fi" else "yyyy-mm-dd"
        row[6].number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    for i, wdt in enumerate((12, 24, 14, 30, 12, 8, 11, 10, 8), 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws2 = wb.create_sheet(W["summary"])
    ws2.append([club["name"], f"{_fi_date(start) if lang == 'fi' else start} – {_fi_date(end) if lang == 'fi' else end}"
                if start != "0000-01-01" else ""])
    ws2["A1"].font = Font(bold=True, size=13)
    ws2.append([])
    ws2.append(W["sum_cols"])
    for c in ws2[3]:
        c.font = Font(bold=True)
    for s in lines:
        ws2.append([kind(s["kind"]), s["what"], s["qty"], s["amount"], s["paid"], s["unpaid"]])
    ws2.append([W["total"], f"{W['count']} {tot['count']}", None, tot["amount"], tot["paid"], tot["unpaid"]])
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    for row in ws2.iter_rows(min_row=4):
        for c in row[3:6]:
            c.number_format = "#,##0.00"
    for i, wdt in enumerate((14, 30, 8, 12, 12, 13), 1):
        ws2.column_dimensions[get_column_letter(i)].width = wdt
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{base}.xlsx"', "Cache-Control": "no-store"})

